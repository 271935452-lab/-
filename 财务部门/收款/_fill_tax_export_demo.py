# -*- coding: utf-8 -*-
"""Fill (V260401) batch invoice import template with demo rows for MVP export."""
import shutil
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent
SRC = ROOT / "(V260401版)批量开票-导入开票模板.xlsx"
DST = ROOT / "批量开票-导入开票-导出示例.xlsx"


def main():
    shutil.copy2(SRC, DST)
    wb = openpyxl.load_workbook(DST)
    ws1 = wb["1-发票基本信息"]
    ws2 = wb["2-发票明细信息"]
    ws3 = wb["3-特定业务信息"]

    # 国内拖车 · 货物运输 · 合开一张
    s1 = "INVREQ20260900801"
    ws1.cell(4, 1).value = s1
    ws1.cell(4, 2).value = "普通发票"
    ws1.cell(4, 3).value = "货物运输"
    ws1.cell(4, 4).value = "是"
    ws1.cell(4, 5).value = "否"
    ws1.cell(4, 6).value = "腾信示例贸易"
    ws1.cell(4, 23).value = "提货/拖车/调拨合开一张"

    ws2.cell(4, 1).value = s1
    ws2.cell(4, 2).value = "*交通运输服务*运输费"
    ws2.cell(4, 5).value = "次"
    ws2.cell(4, 6).value = 1
    ws2.cell(4, 8).value = 2150.00
    ws2.cell(4, 9).value = 0.09

    # 三段车辆明细，同一流水号 → 一张票
    for i, (plate, origin, dest, goods) in enumerate(
        [
            ("粤B·***98", "东莞", "深圳", "家居用品"),
            ("浙G·***16", "东莞", "宁波", "家居用品"),
            ("粤B·***52", "宁波", "宁波", "家居用品"),
        ],
        start=4,
    ):
        ws3.cell(i, 1).value = s1
        ws3.cell(i, 8).value = origin
        ws3.cell(i, 9).value = dest
        ws3.cell(i, 10).value = "公路运输"
        ws3.cell(i, 11).value = plate
        ws3.cell(i, 12).value = goods

    # 国内报关 · 普通开票
    s2 = "INVREQ20260900901"
    ws1.cell(5, 1).value = s2
    ws1.cell(5, 2).value = "普通发票"
    ws1.cell(5, 4).value = "是"
    ws1.cell(5, 5).value = "否"
    ws1.cell(5, 6).value = "深圳市某某贸易有限公司"
    ws1.cell(5, 23).value = "CUSDEC-202609-001、CUSDEC-202609-002"

    ws2.cell(5, 1).value = s2
    ws2.cell(5, 2).value = "*生产生活服务*代理报关费"
    ws2.cell(5, 5).value = "票"
    ws2.cell(5, 6).value = 1
    ws2.cell(5, 8).value = 500.00
    ws2.cell(5, 9).value = 0.06

    wb.save(DST)
    print("ok", DST.name)


if __name__ == "__main__":
    main()
